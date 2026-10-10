# Author: Gorla Poojitha
"""
Build dashboards/category_report.xlsx from data/exports/dashboard_order_items.csv:

    pivot     revenue and orders by category x year, sorted by total revenue,
              with XLOOKUP to the targets sheet and actual vs target %
    targets   illustrative revenue targets for 5 categories (made up for the exercise)
    + a bar chart of the top 10 categories

Run scripts/export_marts.py first.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "data" / "exports" / "dashboard_order_items.csv"
OUTPUT_FILE = PROJECT_ROOT / "dashboards" / "category_report.xlsx"

# Illustrative targets for total revenue across all years (BRL). Not real Olist targets.
TARGETS = {
    "health_beauty": 1_400_000,
    "watches_gifts": 1_200_000,
    "bed_bath_table": 1_300_000,
    "sports_leisure": 1_100_000,
    "computers_accessories": 1_000_000,
}

HEADER_FILL = PatternFill("solid", fgColor="2A78D6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BRL_FORMAT = '"R$"#,##0'


def style_header(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def main() -> None:
    # Exports start with a '# Author' line unless they were made with --plain
    with open(SOURCE, encoding="utf-8") as f:
        skip = 1 if f.readline().startswith("#") else 0
    items = pd.read_csv(SOURCE, parse_dates=["order_date"], skiprows=skip)
    items["order_year"] = items["order_date"].dt.year

    revenue = items.pivot_table(index="category_name", columns="order_year",
                                values="item_total", aggfunc="sum", fill_value=0)
    orders = items.pivot_table(index="category_name", columns="order_year",
                               values="order_id", aggfunc="nunique", fill_value=0)
    years = list(revenue.columns)
    revenue["total"] = revenue.sum(axis=1)
    revenue = revenue.sort_values("total", ascending=False)
    orders = orders.loc[revenue.index]

    wb = Workbook()
    ws = wb.active
    ws.title = "pivot"
    ws["A1"] = "Revenue (item price + freight, BRL) and orders by category and year"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Valid orders only. 2016 and 2018 are partial years. Source: data/exports/dashboard_order_items.csv"
    ws["A3"] = "Prepared by Gorla Poojitha"
    ws["A3"].font = Font(italic=True, color="52514E")

    headers = (["category_name"] + [f"revenue {y}" for y in years] + ["revenue total"]
               + [f"orders {y}" for y in years] + ["target", "actual vs target"])
    header_row = 4
    for col, name in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=name)
    style_header(ws, header_row, len(headers))

    total_col = 2 + len(years)
    target_col = len(headers) - 1
    for i, category in enumerate(revenue.index):
        r = header_row + 1 + i
        ws.cell(row=r, column=1, value=category)
        for j, y in enumerate(years):
            ws.cell(row=r, column=2 + j, value=round(float(revenue.loc[category, y]), 2)).number_format = BRL_FORMAT
            ws.cell(row=r, column=total_col + 1 + j, value=int(orders.loc[category, y]))
        total = ws.cell(row=r, column=total_col,
                        value=f"=SUM({get_column_letter(2)}{r}:{get_column_letter(1 + len(years))}{r})")
        total.number_format = BRL_FORMAT
        # Excel needs the _xlfn prefix for XLOOKUP when the file is written by a library
        target = ws.cell(row=r, column=target_col,
                         value=f'=_xlfn.XLOOKUP(A{r},targets!A:A,targets!B:B,"no target")')
        target.number_format = BRL_FORMAT
        t = get_column_letter(target_col)
        vs = ws.cell(row=r, column=target_col + 1,
                     value=f'=IF(ISNUMBER({t}{r}),{get_column_letter(total_col)}{r}/{t}{r},"")')
        vs.number_format = "0.0%"

    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.column_dimensions["A"].width = 42
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Top 10 categories by revenue"
    chart.y_axis.title = "BRL"
    chart.y_axis.numFmt = '"R$"#,##0'
    chart.legend = None
    data = Reference(ws, min_col=total_col, min_row=header_row, max_row=header_row + 10)
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + 10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.x_axis.scaling.orientation = "maxMin"   # largest at the top
    chart.series[0].graphicalProperties.solidFill = "2A78D6"
    chart.width, chart.height = 22, 11
    ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}4")

    targets = wb.create_sheet("targets")
    targets.append(["category_name", "target_revenue"])
    style_header(targets, 1, 2)
    for category, value in TARGETS.items():
        targets.append([category, value])
        targets.cell(row=targets.max_row, column=2).number_format = BRL_FORMAT
    targets["D1"] = "Illustrative targets for the exercise, not real Olist figures."
    targets.column_dimensions["A"].width = 30
    targets.column_dimensions["B"].width = 18

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(f"Saved {OUTPUT_FILE} ({len(revenue)} categories)")


if __name__ == "__main__":
    main()
